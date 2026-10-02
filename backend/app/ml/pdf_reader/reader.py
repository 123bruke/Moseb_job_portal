"""Document reader: PyMuPDF text per page, OCR fallback for scanned pages, DOCX/TXT support.
Evolved from the old pdf.py (which imported itself and could not run)."""
import io
import logging
from dataclasses import dataclass, field

log = logging.getLogger(__name__)
MIN_CHARS_PER_PAGE = 40


@dataclass
class PageText:
    number: int
    text: str
    ocr: bool = False


@dataclass
class ReadResult:
    pages: list[PageText] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text)

    @property
    def scanned_ratio(self) -> float:
        return sum(p.ocr for p in self.pages) / len(self.pages) if self.pages else 0.0


def sniff(data: bytes) -> str:
    """MIME sniffing by magic bytes (never trust the filename/content-type)."""
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04" and b"word/" in data[:4000] + data[-4000:]:
        return "docx"
    try:
        data[:2000].decode("utf-8")
        if b"\x00" not in data[:2000]:
            return "txt"
    except UnicodeDecodeError:
        pass
    return "unknown"


def page_count(data: bytes, kind: str) -> int:
    if kind == "pdf":
        import fitz
        with fitz.open(stream=data, filetype="pdf") as d:
            return len(d)
    return 1


def _ocr_page(page) -> str:
    try:
        import pytesseract
        from PIL import Image
        pix = page.get_pixmap(dpi=200)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        return pytesseract.image_to_string(img)
    except Exception as e:  # tesseract missing etc. -> degrade gracefully
        log.warning("OCR unavailable: %s", type(e).__name__)
        return ""


def read_pdf(data: bytes, use_layout: bool = True) -> ReadResult:
    import fitz
    from app.ml.yolo.layout import ordered_page_text

    res = ReadResult()
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc, 1):
            text = ordered_page_text(page) if use_layout else page.get_text("text")
            ocr = False
            if len(text.strip()) < MIN_CHARS_PER_PAGE:
                o = _ocr_page(page)
                if len(o.strip()) > len(text.strip()):
                    text, ocr = o, True
            res.pages.append(PageText(i, text.strip(), ocr))
    return res


def read_docx(data: bytes) -> ReadResult:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return ReadResult([PageText(1, "\n".join(parts))])


def read_document(data: bytes, kind: str | None = None) -> ReadResult:
    kind = kind or sniff(data)
    if kind == "pdf":
        return read_pdf(data)
    if kind == "docx":
        return read_docx(data)
    if kind == "txt":
        return ReadResult([PageText(1, data.decode("utf-8", "ignore"))])
    raise ValueError("unsupported document type")
